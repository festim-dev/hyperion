"""Extract experimental permeation curves from the Excel workbook.

Writes one CSV per temperature with columns: time_s, time_h, ppm, flux_mol_m2_s.
"""

from pathlib import Path
import openpyxl

XLSX = Path(__file__).with_name("FLiBe permeability.xlsx")
OUT = Path(__file__).with_name("data")

# (sheet_name, col_time_min, col_time_h, col_ppm, col_flux)
SHEETS = {
    500: ("500C_5mm_coated+swapped+dewar_N", 1, 2, 11, 14),
    550: ("550C_5mm_coated+swapped+dewar_N", 1, 2, 11, 14),
    600: ("600C_5mm_coated+swapped+dewar_N", 1, 2, 11, 14),
    650: ("650C_5mm_coated+swapped+dewar_N", 1, 2, 11, 14),
    700: ("700C_5mm_coated+swapped+dewar_N", 1, 3, 12, 15),
}


def main():
    OUT.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for T_C, (name, c_min, c_h, c_ppm, c_flux) in SHEETS.items():
        ws = wb[name]
        rows = []
        for r in range(2, ws.max_row + 1):
            t_min = ws.cell(r, c_min).value
            t_h = ws.cell(r, c_h).value
            ppm = ws.cell(r, c_ppm).value
            flux = ws.cell(r, c_flux).value
            if t_min is None or flux is None:
                continue
            rows.append((float(t_min) * 60.0, float(t_h), float(ppm), float(flux)))
        path = OUT / f"T{T_C}C.csv"
        with path.open("w") as f:
            f.write("time_s,time_h,ppm,flux_mol_m2_s\n")
            for row in rows:
                f.write(f"{row[0]:.6g},{row[1]:.6g},{row[2]:.6g},{row[3]:.6e}\n")
        print(f"wrote {path}  ({len(rows)} rows)")


if __name__ == "__main__":
    main()
